"""Text extraction for resume/JD files (PDF, DOCX, TXT)."""
from __future__ import annotations

import io

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = ("pdf", "docx", "txt")


class UnsupportedFileType(Exception):
    """Raised when an uploaded file isn't a supported format."""


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a PDF, DOCX, or TXT file's raw bytes."""
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if suffix == "pdf":
        return _extract_pdf(file_bytes)
    if suffix == "docx":
        return _extract_docx(file_bytes)
    if suffix == "txt":
        return file_bytes.decode("utf-8", errors="ignore")

    raise UnsupportedFileType(
        f"Unsupported file type '.{suffix or '?'}'. Please upload a PDF, DOCX, or TXT file."
    )


def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts).strip()
